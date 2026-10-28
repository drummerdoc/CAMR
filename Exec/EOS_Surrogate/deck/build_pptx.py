#!/usr/bin/env python
"""Build the gigaCCS-template .pptx of the EOS-strategies talk from the Slides artifact's files.

    yt_env/bin/python deck/build_pptx.py --template ~/2025-11-12_gigaCCS_MA2_break-out-session.pptx \
        --slides <folder holding project/deck.json and project/slides/*.html> --out deck/eos_strategies_gigaCCS.pptx

The template is opened, all of its slides are removed (its masters, layouts, theme and footers stay), and one
slide per entry of the deck's `order` is added on the template's own layouts: the section-title layout the
other presenters use for the cover, 'Tittel og innhold' for figure slides, 'Tittel og tekst' for the
takeaways.  Titles, captions, bullets, table cells and speaker notes are read from the slide HTML, so a change
in the web deck is a rerun of this script.  Figures are the repo's PNGs, mapped from the artifact's asset ids.
The model diagram is redrawn as native shapes.  Nothing here renders the result: open it in PowerPoint.
"""
import argparse, json, os, re, html as htmlmod
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
FIGDIR = os.path.join(HERE, '..', 'figures')

# artifact asset id -> repo figure
FIGS = {
    'e8c61ee5f32955ed87c2dc1069bf12a6': 'fig1_deployment_cloud.png',
    'e3b4eefb351d2f51d53bd559f71865f8': 'fig3d_combo.png',
    'a3081565aaa8fd9ab26b56d8eb203a2d': 'fig4_accuracy_cost.png',
    '7fc0edb71dc48303101213f6c9a663de': 'fig6_stage_split_linear.png',
    'd43a70aa08391a7178257380a047d988': 'fig11_edge_table.png',
    '15b23bb65efeedcdc9b91b9ebeb66530': 'fig12_dome.png',
    # 2026-10-28: the same figures re-rendered with doubled text (FONT_SCALE=2 FIG_SCALE=1.5)
    '08f9acdafec865899f4c9e5fdf698f67': 'fig1_deployment_cloud.png',
    '5ff5f3afc0349994809a8c6540227f60': 'fig3d_combo.png',
    '029c2bfd946b2e526e48c86d195f2f48': 'fig4_accuracy_cost.png',
    'd117bb9fd4879616e4ff6232948fe1d8': 'fig6_stage_split_linear.png',
    '40ab9c90e8b0bf24e5a23cae82ec892c': 'fig11_edge_table.png',
    '15893f80c2d8672f744fd9c512946a9e': 'fig12_dome.png',
}

INK = RGBColor(0x14, 0x21, 0x3d); INK2 = RGBColor(0x3d, 0x46, 0x54); MUTED = RGBColor(0x6a, 0x71, 0x79)
BLUE = RGBColor(0x2a, 0x78, 0xd6); ORANGE = RGBColor(0xeb, 0x68, 0x34)

# ----------------------------------------------------------------------------- HTML helpers
def strip(s):
    return htmlmod.unescape(re.sub(r'<[^>]+>', '', s)).strip()

def inner(tag, s, attr=''):
    m = re.search(r'<%s\b[^>]*%s[^>]*>(.*?)</%s>' % (tag, attr, tag), s, re.S)
    return m.group(1) if m else ''

def title_of(s):
    return strip(inner('h2', s) or inner('h1', s))

def notes_of(s):
    return strip(inner('aside', s))

def body_paragraphs(s):
    """The non-pinned <p> of a slide (captions), with their inline <b> kept; footers and the eyebrow dropped."""
    out = []
    for m in re.finditer(r'<p\b([^>]*)>(.*?)</p>', s, re.S):
        attrs, body = m.group(1), m.group(2)
        if 'position:absolute' in attrs or 'text-transform:uppercase' in attrs: continue
        out.append(body.strip())
    return out

def fig_of(s):
    m = re.search(r'/_blob/([0-9a-f]{32})', s)
    if not m or m.group(1) not in FIGS: return None
    return os.path.join(FIGDIR, FIGS[m.group(1)])

def runs_from_html(frag):
    """Split an HTML fragment into (text, bold) runs."""
    parts = re.split(r'(<b>.*?</b>)', frag, flags=re.S)
    out = []
    for p in parts:
        if not p: continue
        if p.startswith('<b>'): out.append((strip(p), True))
        else:
            t = strip(p)
            if t: out.append((t, False))
    return out

# ----------------------------------------------------------------------------- pptx helpers
def layout(prs, name):
    for l in prs.slide_masters[0].slide_layouts:
        if l.name == name: return l
    raise KeyError(name)

def remove_all_slides(prs):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        prs.part.drop_rel(sldId.rId); lst.remove(sldId)

def drop_body_placeholders(slide):
    """Remove the empty content/body placeholders a layout clones onto a new slide; keep the title."""
    for sh in list(slide.placeholders):
        if sh.placeholder_format.type is None or sh.placeholder_format.idx != 0:
            sh._element.getparent().remove(sh._element)

def set_title(slide, text, size=None):
    t = slide.shapes.title; t.text = text
    if size:
        for p in t.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(size)

def add_text(slide, left, top, width, height, paragraphs, size=12, color=INK2, anchor=MSO_ANCHOR.TOP, space_after=4):
    """paragraphs: list of HTML fragments (with <b>) or (fragment, dict) tuples; dict may set size, bold, color, bullet."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for item in paragraphs:
        frag, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(opt.get('space_after', space_after))
        if opt.get('align') == 'right': p.alignment = PP_ALIGN.RIGHT
        for text, bold in runs_from_html(frag):
            r = p.add_run(); r.text = text
            r.font.size = Pt(opt.get('size', size)); r.font.bold = bold or opt.get('bold', False)
            r.font.color.rgb = opt.get('color', color)
    return tb

def add_figure(slide, path, left, top, max_w, max_h):
    """Add a picture scaled to fit (max_w x max_h) inches, centred in that box."""
    from PIL import Image
    w_px, h_px = Image.open(path).size
    scale = min(max_w / w_px, max_h / h_px); w, h = w_px * scale, h_px * scale
    return slide.shapes.add_picture(path, Inches(left + (max_w - w) / 2), Inches(top + (max_h - h) / 2), Inches(w), Inches(h))

def set_notes(slide, text):
    if text: slide.notes_slide.notes_text_frame.text = text

# ----------------------------------------------------------------------------- slide recipes
# Content area of 'Tittel og innhold': x 0.8..12.6 in, y 2.5..6.8 in; the title band ends at 1.5 in.
X0, W = 0.8, 11.8

def slide_cover(prs, s, author_line):
    sl = prs.slides.add_slide(layout(prs, '1_title slide_v1'))
    tb = add_text(sl, 0.2, 5.9, 12.9, 1.2, [(title_of(s), {'size': 28, 'bold': True, 'color': INK}), (author_line, {'size': 18, 'color': INK2})], space_after=2)
    set_notes(sl, notes_of(s)); return sl

def slide_figure(prs, s, fig_h=2.9, cap_size=16):
    sl = prs.slides.add_slide(layout(prs, 'Tittel og innhold')); set_title(sl, title_of(s)); drop_body_placeholders(sl)
    fig = fig_of(s)
    if fig: add_figure(sl, fig, X0, 1.7, W, fig_h)
    add_text(sl, X0, 1.75 + fig_h, W, 6.75 - (1.75 + fig_h), body_paragraphs(s), size=cap_size)
    set_notes(sl, notes_of(s)); return sl

def slide_cloud(prs, s):
    sl = prs.slides.add_slide(layout(prs, 'Tittel og innhold')); set_title(sl, title_of(s)); drop_body_placeholders(sl)
    paras = body_paragraphs(s)                       # [callout, third, no-root, corridor]
    callout, rest = paras[0], paras[1:]
    add_text(sl, X0, 1.7, 3.3, 3.3, [('<b>“Harvest”</b>', {'size': 16, 'color': INK}), (callout, {'size': 14})])
    add_figure(sl, fig_of(s), X0 + 3.4, 1.7, W - 3.4, 3.3)
    cw = (W - 0.4) / 3
    for i, para in enumerate(rest[:3]):
        add_text(sl, X0 + i * (cw + 0.2), 5.1, cw, 1.7, [para], size=14)
    set_notes(sl, notes_of(s)); return sl

def slide_dome(prs, s):
    sl = prs.slides.add_slide(layout(prs, 'Tittel og innhold')); set_title(sl, title_of(s)); drop_body_placeholders(sl)
    add_figure(sl, fig_of(s), X0, 1.7, 6.9, 4.9)
    items = []
    for h, p in re.findall(r'<h3\b[^>]*>(.*?)</h3>\s*<p\b[^>]*>(.*?)</p>', s, re.S):
        items.append((strip(h), {'size': 16, 'bold': True, 'color': INK, 'space_after': 1})); items.append((p, {'size': 14, 'space_after': 8}))
    add_text(sl, X0 + 7.1, 1.6, W - 7.1, 5.2, items)
    set_notes(sl, notes_of(s)); return sl

def slide_whyml(prs, s):
    sl = prs.slides.add_slide(layout(prs, 'Tittel og innhold')); set_title(sl, title_of(s)); drop_body_placeholders(sl)
    paras = body_paragraphs(s)
    add_text(sl, X0, 1.65, W, 1.1, [paras[0]], size=16)
    rows = [[c for c in re.findall(r'<t[hd]\b[^>]*>(.*?)</t[hd]>', r, re.S)] for r in re.findall(r'<tr\b[^>]*>(.*?)</tr>', s, re.S)]
    nr, nc = len(rows), len(rows[0])
    tbl = sl.shapes.add_table(nr, nc, Inches(X0), Inches(2.8), Inches(W), Inches(0.48 * nr)).table
    for ci, wf in enumerate((0.40, 0.30, 0.30)): tbl.columns[ci].width = Inches(W * wf)
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = tbl.cell(ri, ci); c.text = ''
            p = c.text_frame.paragraphs[0]
            if ci > 0: p.alignment = PP_ALIGN.RIGHT
            for text, bold in runs_from_html(cell):
                r = p.add_run(); r.text = text; r.font.size = Pt(15); r.font.bold = bold or ri == 0
                if 'eb6834' in cell: r.font.color.rgb = ORANGE
    add_text(sl, X0, 2.95 + 0.5 * nr, W, 6.75 - (2.95 + 0.5 * nr), paras[1:], size=15)
    set_notes(sl, notes_of(s)); return sl

def slide_takeaways(prs, s):
    sl = prs.slides.add_slide(layout(prs, 'Tittel og tekst')); set_title(sl, title_of(s), size=26)
    body = [sh for sh in sl.placeholders if sh.placeholder_format.idx != 0][0]
    tf = body.text_frame; tf.word_wrap = True
    items = re.findall(r'<li\b[^>]*>(.*?)</li>', s, re.S)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(6)
        for text, bold in runs_from_html(it):          # the template's body placeholder supplies the bullets
            r = p.add_run(); r.text = text; r.font.size = Pt(18); r.font.bold = bold
    set_notes(sl, notes_of(s)); return sl

def slide_model(prs, s):
    """The three-box diagram, redrawn with native shapes from the SVG's text."""
    sl = prs.slides.add_slide(layout(prs, 'Tittel og innhold')); set_title(sl, title_of(s)); drop_body_placeholders(sl)
    svg = inner('svg', s)
    texts = [(float(x), float(y), strip(t), 'font-weight="600"' in a, re.search(r'fill="(#[0-9a-f]{6})"', a).group(1))
             for a, x, y, t in [(m.group(1), m.group(2), m.group(3), m.group(4)) for m in re.finditer(r'<text\b([^>]*x="([\d.]+)"[^>]*y="([\d.]+)"[^>]*)>(.*?)</text>', svg, re.S)]]
    rects = [(float(x), float(w), col) for x, w, col in re.findall(r'<rect x="([\d.]+)" y="60" width="([\d.]+)"[^>]*stroke="(#[0-9a-f]{6})"', svg)]
    sx = W / 1664.0                                         # svg px -> inches
    top, box_h = 1.75, 2.85
    for x, w, col in rects:
        shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(X0 + x * sx), Inches(top), Inches(w * sx), Inches(box_h))
        shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(0xff, 0xf4, 0xee) if col == '#eb6834' else RGBColor(0xff, 0xff, 0xff)
        shp.line.color.rgb = RGBColor.from_string(col[1:]); shp.line.width = Pt(1.5); shp.shadow.inherit = False
        shp.adjustments[0] = 0.06
        inside = [t for t in texts if x <= t[0] < x + w and 60 <= t[1] <= 390]
        inside.sort(key=lambda t: t[1])
        paras = []
        for tx, ty, txt, bold, fill in inside:
            paras.append((htmlmod.escape(txt), {'size': 16 if bold else 13, 'bold': bold, 'color': RGBColor.from_string(fill[1:]), 'space_after': 2}))
        add_text(sl, X0 + x * sx + 0.08, top + 0.08, w * sx - 0.16, box_h - 0.16, paras)
    for i in range(len(rects) - 1):
        x_end = (rects[i][0] + rects[i][1]) * sx
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(X0 + x_end + 0.02), Inches(top + box_h / 2), Inches(X0 + rects[i + 1][0] * sx - 0.02), Inches(top + box_h / 2))
        c.line.color.rgb = MUTED; c.line.width = Pt(2)
        c.line._get_or_add_ln().append(c.line._get_or_add_ln().makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd', {'type': 'triangle'}))
    below = [t for t in texts if t[1] > 400]; below.sort(key=lambda t: t[1])
    head = [t for t in texts if t[1] < 60]
    add_text(sl, X0, 1.4, W, 0.35, [(htmlmod.escape(t[2]), {'size': 15, 'bold': True, 'color': INK}) for t in head])
    add_text(sl, X0, top + box_h + 0.15, W, 6.75 - (top + box_h + 0.15),
             [(htmlmod.escape(t[2]), {'size': 15 if t[3] else 14, 'bold': t[3], 'color': RGBColor.from_string(t[4][1:]), 'space_after': 3}) for t in below])
    set_notes(sl, notes_of(s)); return sl

RECIPES = {'cover': slide_cover, 'model': slide_model, 'dome': slide_dome, 'cloud': slide_cloud, 'whyml': slide_whyml, 'takeaways': slide_takeaways}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--template', required=True); ap.add_argument('--slides', required=True); ap.add_argument('--out', required=True)
    ap.add_argument('--author', default='Marc Day, CAMR · Pelanti–Shyue six-equation model · October 2026')
    a = ap.parse_args()
    prs = Presentation(a.template); remove_all_slides(prs)
    deck = json.load(open(os.path.join(a.slides, 'project', 'deck.json')))
    for sid in deck['order']:
        s = open(os.path.join(a.slides, 'project', 'slides', sid + '.html'), encoding='utf-8').read()
        fn = RECIPES.get(sid)
        if fn is slide_cover: fn(prs, s, a.author)
        elif fn: fn(prs, s)
        else: slide_figure(prs, s)
        print(f'  {sid:10s} {title_of(s)[:70]}')
    prs.save(a.out); print('wrote', a.out, f'({os.path.getsize(a.out) / 1e6:.1f} MB)')

if __name__ == '__main__':
    main()
